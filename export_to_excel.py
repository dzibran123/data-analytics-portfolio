from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


PAID_FILE = Path("cleaned_sales_paid.csv")
CANCELLED_FILE = Path("sales_cancelled.csv")
REFUNDED_FILE = Path("sales_refunded.csv")
OUTPUT_FILE = Path("Laporan_Penjualan_Maret_2026.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="168AAD")
SECTION_FILL = PatternFill("solid", fgColor="D9EEF2")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=16, bold=True, color="24313A")
SECTION_FONT = Font(bold=True, color="24313A")
CURRENCY_FORMAT = '"Rp" #,##0'
INTEGER_FORMAT = "#,##0"
THIN_BORDER = Border(
    left=Side(style="thin", color="D9E1E5"),
    right=Side(style="thin", color="D9E1E5"),
    top=Side(style="thin", color="D9E1E5"),
    bottom=Side(style="thin", color="D9E1E5"),
)


def read_input_files() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    paid = pd.read_csv(PAID_FILE)
    cancelled = pd.read_csv(CANCELLED_FILE)
    refunded = pd.read_csv(REFUNDED_FILE)

    for sales in (paid, cancelled, refunded):
        sales["total_harga"] = pd.to_numeric(
            sales["total_harga"], errors="coerce"
        ).abs()
        sales["jumlah_beli"] = pd.to_numeric(
            sales["jumlah_beli"], errors="coerce"
        )

    return paid, cancelled, refunded


def apply_borders_and_autofit(worksheet) -> None:
    for row in worksheet.iter_rows(
        min_row=1, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column
    ):
        for cell in row:
            if cell.value is not None:
                cell.border = THIN_BORDER

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 45)


def format_numeric_columns(worksheet) -> None:
    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }
    for column_name in ("harga_satuan", "total_harga", "Total Omzet"):
        column = headers.get(column_name)
        if column is not None:
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row=row, column=column).number_format = CURRENCY_FORMAT
    quantity_column = headers.get("jumlah_beli")
    if quantity_column is not None:
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row, column=quantity_column).number_format = INTEGER_FORMAT


def style_data_sheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    format_numeric_columns(worksheet)
    apply_borders_and_autofit(worksheet)


def style_summary_sheet(worksheet) -> None:
    worksheet.freeze_panes = "A4"
    worksheet.sheet_view.showGridLines = False
    worksheet["A1"].font = TITLE_FONT
    for row in (3, 11):
        for cell in worksheet[row]:
            if cell.value is not None:
                cell.fill = SECTION_FILL
                cell.font = SECTION_FONT

    for row in (4, 12):
        for cell in worksheet[row]:
            if cell.value is not None:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center")

    for row in range(5, 8):
        worksheet.cell(row=row, column=2).number_format = CURRENCY_FORMAT
    for row in range(13, worksheet.max_row + 1):
        worksheet.cell(row=row, column=5).number_format = CURRENCY_FORMAT
        worksheet.cell(row=row, column=8).number_format = CURRENCY_FORMAT
    apply_borders_and_autofit(worksheet)


def build_summary(
    paid: pd.DataFrame,
    cancelled: pd.DataFrame,
    refunded: pd.DataFrame,
) -> pd.DataFrame:
    problem_transactions = pd.concat([cancelled, refunded], ignore_index=True)
    problem_product = "Tidak ada"
    if not problem_transactions.empty:
        problem_product = problem_transactions["nama_produk"].value_counts().idxmax()

    return pd.DataFrame(
        {
            "Metrik": [
                "Total Revenue Bersih",
                "Total Potensi Batal",
                "Total Kerugian Refund",
                "Produk Paling Sering Bermasalah",
            ],
            "Nilai": [
                paid["total_harga"].sum(),
                cancelled["total_harga"].sum(),
                refunded["total_harga"].sum(),
                problem_product,
            ],
        }
    )


def main() -> None:
    paid, cancelled, refunded = read_input_files()
    category_summary = (
        paid.groupby("kategori", as_index=False)["total_harga"]
        .sum()
        .rename(columns={"total_harga": "Total Omzet"})
        .sort_values("Total Omzet", ascending=False)
    )
    city_summary = (
        paid.groupby("kota_pembeli", as_index=False)["total_harga"]
        .sum()
        .rename(columns={"total_harga": "Total Omzet"})
        .sort_values("Total Omzet", ascending=False)
    )

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        summary = build_summary(paid, cancelled, refunded)
        summary.to_excel(writer, sheet_name="Ringkasan KPI", index=False, startrow=3)
        category_summary.to_excel(
            writer, sheet_name="Ringkasan KPI", index=False, startrow=11, startcol=3
        )
        city_summary.to_excel(
            writer, sheet_name="Ringkasan KPI", index=False, startrow=11, startcol=6
        )

        paid.to_excel(writer, sheet_name="Transaksi Sukses (PAID)", index=False)
        cancelled.to_excel(writer, sheet_name="Pesanan Batal (CANCEL)", index=False)
        refunded.to_excel(writer, sheet_name="Retur (REFUND)", index=False)

        workbook = writer.book
        summary_sheet = workbook["Ringkasan KPI"]
        summary_sheet["A1"] = "Laporan Penjualan Maret 2026"
        summary_sheet["A3"] = "Ringkasan KPI"
        summary_sheet["D11"] = "Omzet per Kategori Produk"
        summary_sheet["G11"] = "Omzet per Kota Pembeli"

        for worksheet in workbook.worksheets:
            if worksheet.title == "Ringkasan KPI":
                style_summary_sheet(worksheet)
            else:
                style_data_sheet(worksheet)

    print(f"Laporan Excel berhasil disimpan ke: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()